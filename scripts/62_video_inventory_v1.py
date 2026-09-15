#!/usr/bin/env python3
"""Step V1: 培训视频盘点（deterministic，无 LLM）。

目标：
- 把 video_learning 项目的 22 个培训视频（metadata.csv / videos.json / SRT）
  收敛为一份带 OSS 外链与溯源的 inventory。
- 与 Step 13.1 的 21 份正式文档（衍景CRM*.md，已确认是培训 ASR 转写）
  做文本包含关系比对，判定每个视频是否已有文档覆盖。
- 只盘点、不判定知识结构，不产出 KB 条目。

输入：
- VIDEO_ROOT/learning_center/metadata.csv
- VIDEO_ROOT/learning_center/output/videos.json
- VIDEO_ROOT/learning_center/transcripts/*.srt
- output/formal_docs_inventory_v1.jsonl（classification=formal_document 的 21 份）
- data/raw/markdowns/<filename>

输出：
- output/video_inventory_v1.jsonl
- output/video_inventory_v1.xlsx（summary / videos / doc_matches /
  unmatched_videos / unmatched_formal_docs）

可选：
- --check-urls 对 22 个视频与字幕 URL 逐个发 HEAD 请求验证公网可达。
"""

import argparse
import csv
import difflib
import json
import re
import sys
import urllib.request
from pathlib import Path

from openpyxl import Workbook

VIDEO_ROOT = Path("/Users/qiyue/Desktop/video_learning")
KB_ROOT = Path(__file__).resolve().parent.parent
OUT_JSONL = KB_ROOT / "output" / "video_inventory_v1.jsonl"
OUT_XLSX = KB_ROOT / "output" / "video_inventory_v1.xlsx"
FORMAL_INVENTORY = KB_ROOT / "output" / "formal_docs_inventory_v1.jsonl"
MARKDOWN_DIR = KB_ROOT / "data" / "raw" / "markdowns"

OSS_BASE = "https://aierka-car.oss-cn-hangzhou.aliyuncs.com/"

# shingle 参数与阈值
SHINGLE_N = 25
SHINGLE_STRIDE = 5
FULL_MATCH = 0.90
PARTIAL_MATCH = 0.30

NORMALIZE_RE = re.compile(r"[\s，。、；：？！“”‘’（）,.:;?!()\-—…]")


def norm_text(s: str) -> str:
    return NORMALIZE_RE.sub("", s)


def parse_srt(path: Path):
    raw = path.read_text(encoding="utf-8")
    segments = []
    for block in re.split(r"\n\s*\n", raw.strip()):
        lines = [l for l in block.split("\n") if l.strip()]
        if not lines:
            continue
        ts = None
        text_rows = []
        for line in lines:
            if "-->" in line:
                ts = line.strip()
            elif not line.strip().isdigit():
                text_rows.append(line.strip())
        if ts and text_rows:
            segments.append({"time": ts, "text": "".join(text_rows)})
    return segments


def shingles(s: str, n: int = SHINGLE_N, stride: int = SHINGLE_STRIDE):
    if len(s) <= n:
        return {s} if s else set()
    return {s[i:i + n] for i in range(0, len(s) - n + 1, stride)}


def coverage(src: str, dst_shingles: set) -> float:
    """src 的 shingle 有多少比例出现在 dst 中。"""
    src_sh = shingles(src)
    if not src_sh:
        return 0.0
    return len(src_sh & dst_shingles) / len(src_sh)


def head_url(url: str, timeout: int = 15):
    req = urllib.request.Request(url, method="HEAD")
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.status, resp.headers.get("Content-Length")
    except Exception as e:  # noqa: BLE001 - 记录任何网络错误
        return None, str(e)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--check-urls", action="store_true",
                    help="HEAD 校验每个视频/字幕 URL 的公网可达性")
    args = ap.parse_args()

    errors = []
    warnings = []

    # ---- 1. 读 metadata.csv ----
    meta_path = VIDEO_ROOT / "learning_center" / "metadata.csv"
    with open(meta_path, encoding="utf-8") as f:
        meta_rows = list(csv.DictReader(f))
    meta_by_id = {}
    for r in meta_rows:
        vid = Path(r["filename"]).stem  # video-001.mov -> video-001
        meta_by_id[vid] = r

    # ---- 2. 读 videos.json ----
    schema_path = VIDEO_ROOT / "learning_center" / "output" / "videos.json"
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    schema_by_id = {v["id"]: v for v in schema.get("videos", [])}

    if set(meta_by_id) != set(schema_by_id):
        errors.append(
            f"metadata.csv 与 videos.json 的 id 集合不一致: "
            f"meta-only={sorted(set(meta_by_id) - set(schema_by_id))} "
            f"schema-only={sorted(set(schema_by_id) - set(meta_by_id))}"
        )

    # ---- 3. 解析 SRT 并组装记录 ----
    srt_dir = VIDEO_ROOT / "learning_center" / "transcripts"
    videos = []
    for vid in sorted(set(meta_by_id) | set(schema_by_id)):
        meta = meta_by_id.get(vid, {})
        sch = schema_by_id.get(vid, {})
        srt_path = srt_dir / f"{vid}.srt"
        if not srt_path.exists():
            errors.append(f"{vid}: 缺少字幕文件 {srt_path.name}")
            segments, transcript_text = [], ""
        else:
            segments = parse_srt(srt_path)
            if not segments:
                errors.append(f"{vid}: SRT 解析为空")
            transcript_text = norm_text("".join(s["text"] for s in segments))

        video_key = sch.get("videoUrl") or meta.get("video_compressed_url", "")
        transcript_key = sch.get("transcriptUrl") or meta.get("transcript_url", "")
        for key, label in ((video_key, "video"), (transcript_key, "transcript")):
            if key and not key.startswith("help/"):
                warnings.append(f"{vid}: {label} OSS key 未按 help/ 前缀: {key}")

        videos.append({
            "video_id": vid,
            "source_filename": meta.get("filename", ""),
            "title": meta.get("title") or sch.get("title", ""),
            "summary": meta.get("summary") or sch.get("summary", ""),
            "category": meta.get("category") or sch.get("category", ""),
            "sort": meta.get("sort", ""),
            "duration_sec": float(meta["duration"]) if meta.get("duration") else sch.get("duration"),
            "video_url": OSS_BASE + video_key if video_key else "",
            "transcript_url": OSS_BASE + transcript_key if transcript_key else "",
            "srt_file": srt_path.name if srt_path.exists() else "",
            "srt_segments": len(segments),
            "transcript_chars_normalized": len(transcript_text),
            "_transcript_text": transcript_text,
        })

    # ---- 4. 读 21 份正式文档原文 ----
    formal_rows = [json.loads(l) for l in FORMAL_INVENTORY.read_text(encoding="utf-8").splitlines() if l.strip()]
    formal_docs = {}
    for r in formal_rows:
        if r.get("classification") != "formal_document":
            continue
        p = MARKDOWN_DIR / r["filename"]
        if not p.exists():
            errors.append(f"正式文档缺原文: {r['filename']}")
            continue
        formal_docs[r["document_id"]] = {
            "document_id": r["document_id"],
            "filename": r["filename"],
            "text": norm_text(p.read_text(encoding="utf-8")),
        }
    doc_shingle_cache = {d: shingles(v["text"]) for d, v in formal_docs.items()}

    # ---- 5. 视频 ↔ 正式文档 双向包含比对 ----
    # 第一 pass：字符 shingle 包含率（快，但对个别字符差异造成的整体错位敏感，
    # 例如 ASR 把"衍景CRM"转成"Cm"会使全部 25 字窗口错开）。
    # 第二 pass：对 shingle 判 none/partial 的视频用 difflib.SequenceMatcher
    # 兜底（先用 quick_ratio 预筛再算 ratio），取最优候选。
    doc_matches = []
    for v in videos:
        best = None
        shingle_scores = {}
        for did, dsh in doc_shingle_cache.items():
            vid_in_doc = coverage(v["_transcript_text"], dsh)
            doc_in_vid = coverage(formal_docs[did]["text"], shingles(v["_transcript_text"]))
            shingle_scores[did] = (vid_in_doc, doc_in_vid)
            if best is None or vid_in_doc > best["video_in_doc"]:
                best = {
                    "video_id": v["video_id"],
                    "document_id": did,
                    "doc_filename": formal_docs[did]["filename"],
                    "video_in_doc": round(vid_in_doc, 4),
                    "doc_in_video": round(doc_in_vid, 4),
                }

        best_by_shingles = best["video_in_doc"]
        ratio_used = False
        if best_by_shingles < FULL_MATCH:
            sm_ratio = 0.0
            sm_best_did = ""
            vt = v["_transcript_text"]
            for did in formal_docs:
                dtext = formal_docs[did]["text"]
                sm = difflib.SequenceMatcher(None, vt, dtext)
                if sm.real_quick_ratio() < FULL_MATCH and sm.quick_ratio() < 0.5:
                    continue
                r = sm.ratio()
                if r > sm_ratio:
                    sm_ratio, sm_best_did = r, did
            if sm_best_did and sm_ratio > best_by_shingles:
                best = {
                    "video_id": v["video_id"],
                    "document_id": sm_best_did,
                    "doc_filename": formal_docs[sm_best_did]["filename"],
                    "video_in_doc": round(sm_ratio, 4),
                    "doc_in_video": round(sm_ratio, 4),
                    "method": "difflib_ratio",
                }
                ratio_used = True

        score = best["video_in_doc"]
        if score >= FULL_MATCH:
            best["match_type"] = "full"
        elif score >= PARTIAL_MATCH:
            best["match_type"] = "partial"
        else:
            best["match_type"] = "none"
        if not ratio_used:
            best["method"] = "shingle_coverage"
        doc_matches.append(best)
        v["doc_match"] = best["match_type"]
        v["matched_document_id"] = best["document_id"] if best["match_type"] != "none" else ""

    matched_doc_ids = {m["document_id"] for m in doc_matches if m["match_type"] in ("full", "partial")}
    unmatched_docs = [
        {"document_id": d["document_id"], "filename": d["filename"], "chars": len(d["text"])}
        for d in formal_docs.values() if d["document_id"] not in matched_doc_ids
    ]
    unmatched_videos = [
        {"video_id": v["video_id"], "title": v["title"], "transcript_chars": v["transcript_chars_normalized"]}
        for v in videos if v["doc_match"] == "none"
    ]

    # ---- 6. 可选 URL 活性校验 ----
    url_checks = []
    if args.check_urls:
        for v in videos:
            s1, len1 = head_url(v["video_url"])
            s2, len2 = head_url(v["transcript_url"])
            url_checks.append({
                "video_id": v["video_id"],
                "video_status": s1, "video_bytes": len1,
                "transcript_status": s2,
            })
            if s1 != 200:
                errors.append(f"{v['video_id']}: 视频 URL 不可达 ({s1})")
            if s2 != 200:
                errors.append(f"{v['video_id']}: 字幕 URL 不可达 ({s2})")

    # ---- 7. 汇总守卫 ----
    n_full = sum(1 for m in doc_matches if m["match_type"] == "full")
    n_partial = sum(1 for m in doc_matches if m["match_type"] == "partial")
    n_none = sum(1 for m in doc_matches if m["match_type"] == "none")
    summary = {
        "videos_total": len(videos),
        "srt_segments_total": sum(v["srt_segments"] for v in videos),
        "transcript_chars_total": sum(v["transcript_chars_normalized"] for v in videos),
        "formal_docs_total": len(formal_docs),
        "doc_match_full": n_full,
        "doc_match_partial": n_partial,
        "doc_match_none": n_none,
        "unmatched_formal_docs": len(unmatched_docs),
        "url_check": "done" if args.check_urls else "skipped",
        "errors": len(errors),
        "warnings": len(warnings),
    }

    # ---- 8. 写 jsonl ----
    OUT_JSONL.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT_JSONL, "w", encoding="utf-8") as f:
        for v in videos:
            rec = {k: val for k, val in v.items() if k != "_transcript_text"}
            rec["transcript_text_normalized"] = v["_transcript_text"]
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")

    # ---- 9. 写 xlsx ----
    wb = Workbook()

    ws = wb.active
    ws.title = "summary"
    ws.append(["key", "value"])
    for k, val in summary.items():
        ws.append([k, val])
    ws.append([])
    ws.append(["errors"])
    for e in errors:
        ws.append([e])
    ws.append([])
    ws.append(["warnings"])
    for w in warnings:
        ws.append([w])

    ws = wb.create_sheet("videos")
    cols = ["video_id", "source_filename", "title", "category", "summary",
            "duration_sec", "sort", "video_url", "transcript_url",
            "srt_segments", "transcript_chars_normalized", "doc_match",
            "matched_document_id"]
    ws.append(cols)
    for v in videos:
        ws.append([v[c] for c in cols])

    ws = wb.create_sheet("doc_matches")
    ws.append(["video_id", "match_type", "method", "document_id", "doc_filename",
               "video_in_doc", "doc_in_video"])
    for m in doc_matches:
        ws.append([m["video_id"], m["match_type"], m.get("method", ""),
                   m["document_id"], m["doc_filename"], m["video_in_doc"],
                   m["doc_in_video"]])

    ws = wb.create_sheet("unmatched_videos")
    ws.append(["video_id", "title", "transcript_chars"])
    for r in unmatched_videos:
        ws.append([r["video_id"], r["title"], r["transcript_chars"]])

    ws = wb.create_sheet("unmatched_formal_docs")
    ws.append(["document_id", "filename", "chars"])
    for r in unmatched_docs:
        ws.append([r["document_id"], r["filename"], r["chars"]])

    if url_checks:
        ws = wb.create_sheet("url_checks")
        ws.append(["video_id", "video_status", "video_bytes", "transcript_status"])
        for r in url_checks:
            ws.append([r["video_id"], r["video_status"], r["video_bytes"],
                       r["transcript_status"]])

    wb.save(OUT_XLSX)

    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"\nwrote: {OUT_JSONL}")
    print(f"wrote: {OUT_XLSX}")
    if errors:
        print("\nERRORS:")
        for e in errors:
            print(" -", e)
        sys.exit(1)


if __name__ == "__main__":
    main()
